"""Stage 5: Excel export."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from xpath_gen.discovery.merge import discovery_stats
from xpath_gen.models import RawElement, ValidatedLocator


HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def _status_cell(ws, row: int, col: int, passed: bool) -> None:
    cell = ws.cell(row=row, column=col, value="PASS" if passed else "FAIL")
    cell.fill = PASS_FILL if passed else FAIL_FILL
    cell.alignment = Alignment(horizontal="center")


def export_to_excel(
    locators: list[ValidatedLocator],
    *,
    url: str,
    output_path: Path,
    raw_elements: list[RawElement] | None = None,
) -> Path:
    """Write locators and summary to an Excel workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Locators"

    headers = [
        "Element Name",
        "Section",
        "Tag",
        "Text",
        "XPath",
        "Confidence",
        "Chromium",
        "Firefox",
        "WebKit",
        "Overall",
        "Discovery Source",
        "Notes",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    browser_cols = {"chromium": 7, "firefox": 8, "webkit": 9}

    for row_idx, loc in enumerate(locators, 2):
        ws.cell(row=row_idx, column=1, value=loc.element_name)
        ws.cell(row=row_idx, column=2, value=loc.section.value)
        ws.cell(row=row_idx, column=3, value=loc.tag)
        ws.cell(row=row_idx, column=4, value=loc.text[:100])
        ws.cell(row=row_idx, column=5, value=loc.xpath)
        ws.cell(row=row_idx, column=6, value=loc.confidence.value)

        browser_pass = {r.browser: r.passed for r in loc.browser_results}
        for browser, col in browser_cols.items():
            if browser in browser_pass:
                _status_cell(ws, row_idx, col, browser_pass[browser])
            else:
                ws.cell(row=row_idx, column=col, value="N/A")

        overall = "PASS" if loc.overall_pass else "FAIL"
        overall_cell = ws.cell(row=row_idx, column=10, value=overall)
        overall_cell.fill = PASS_FILL if loc.overall_pass else FAIL_FILL

        ws.cell(row=row_idx, column=11, value=loc.discovery_source.value)
        notes = loc.repair_note or loc.rationale
        ws.cell(row=row_idx, column=12, value=notes[:200])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # Summary sheet
    summary = wb.create_sheet("Summary")
    total = len(locators)
    passed = sum(1 for l in locators if l.overall_pass)
    pass_rate = (passed / total * 100) if total else 0.0

    static_count, discovered_count, boost_pct = (0, 0, 0.0)
    if raw_elements:
        static_count, discovered_count, boost_pct = discovery_stats(raw_elements)

    summary_data = [
        ("URL", url),
        ("Generated At", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("Total Elements", total),
        ("Passed Validation", passed),
        ("Pass Rate (%)", round(pass_rate, 1)),
        ("Static Elements", static_count),
        ("Discovered Elements", discovered_count),
        ("Discovery Boost (%)", boost_pct),
    ]

    for row_idx, (label, value) in enumerate(summary_data, 1):
        summary.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        summary.cell(row=row_idx, column=2, value=value)

    # Section breakdown
    section_counts: dict[str, int] = {}
    for loc in locators:
        section_counts[loc.section.value] = section_counts.get(loc.section.value, 0) + 1

    summary.cell(row=11, column=1, value="Section Breakdown").font = Font(bold=True)
    row = 12
    for section, count in sorted(section_counts.items()):
        summary.cell(row=row, column=1, value=section)
        summary.cell(row=row, column=2, value=count)
        row += 1

    wb.save(output_path)
    return output_path
