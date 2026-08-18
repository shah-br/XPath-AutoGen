"""Tests for Excel export."""

from pathlib import Path

from openpyxl import load_workbook

from xpath_gen.models import (
    BrowserValidation,
    Confidence,
    DiscoverySource,
    Section,
    ValidationStatus,
    ValidatedLocator,
)
from xpath_gen.pipeline.export import export_to_excel


def test_export_creates_workbook(tmp_path: Path):
    locators = [
        ValidatedLocator(
            element_id="el_1",
            element_name="LoginButton",
            xpath="//*[@id='login']",
            confidence=Confidence.HIGH,
            rationale="Unique id",
            section=Section.FORM,
            tag="button",
            text="Login",
            discovery_source=DiscoverySource.STATIC,
            browser_results=[
                BrowserValidation(browser="chromium", passed=True, element_count=1),
                BrowserValidation(browser="firefox", passed=True, element_count=1),
            ],
            overall_status=ValidationStatus.VALID,
        )
    ]
    out = tmp_path / "test.xlsx"
    export_to_excel(locators, url="https://example.com", output_path=out)

    assert out.exists()
    wb = load_workbook(out)
    assert "Locators" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert wb["Locators"].cell(2, 1).value == "LoginButton"
    assert wb["Summary"].cell(3, 2).value == 1
